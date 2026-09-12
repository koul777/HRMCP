"""Versioned Excel replacement: source deltas, dependency retirement and reconciliation.

Only the newly-created candidate is writable. Retired source records are archived
verbatim; no raw KSA text is updated. Human decisions touching a delta block it.
"""
from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import stat
import tempfile
from contextlib import closing
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from .db import connect, initialize_database, now_utc
from .api_refresh_builder import file_sha256, raw_ksa_sha256, trusted_review_status_identity_digest
from .builder_authorization import (
    BuilderAuthorizationError, BuilderOperationContext, require_builder_context,
)
from .ontology_refresh_builder import _run_pipeline, _sqlite_online_snapshot
from .preprocess_excel import HEADER_ALIASES, Normalizer, build_header_map, get
from .sqlite_diagnostics import is_dbstat_table

FIELDS = tuple(HEADER_ALIASES)
TRUSTED = "'human_reviewed','accepted','reviewed'"


def _q(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _payload(values: dict[str, str]) -> str:
    return json.dumps([values[field] for field in FIELDS], ensure_ascii=False, separators=(',', ':'))


def _fingerprints(conn: sqlite3.Connection, table: str) -> dict[str, str]:
    result: dict[str, str] = {}
    current = None
    digest = hashlib.sha256()
    # A set ignores duplicate imports as well as sheet/row ordering.
    for unit, value in conn.execute(f'SELECT DISTINCT unit, digest FROM {_q(table)} ORDER BY unit, digest'):
        if current != unit:
            if current is not None:
                result[current] = digest.hexdigest()
            current, digest = unit, hashlib.sha256()
        digest.update(value.encode('ascii'))
    if current is not None:
        result[current] = digest.hexdigest()
    return result


def _stage(excel: Path, baseline: Path, staging: Path,
           progress: Callable[[str | dict[str, Any]], None] | None = None) -> tuple[dict, sqlite3.Connection]:
    emit = progress or (lambda message: None)
    conn = sqlite3.connect(staging)
    try:
        conn.executescript('CREATE TABLE incoming(unit TEXT,digest TEXT,payload TEXT,sheet TEXT,row_number INTEGER);'
                           'CREATE TABLE previous(unit TEXT,digest TEXT);')
        with closing(sqlite3.connect(baseline.as_uri() + '?mode=ro', uri=True)) as old:
            old.row_factory = sqlite3.Row
            previous_majors = {row[0] for row in old.execute('SELECT DISTINCT major_code FROM raw_excel_rows')}
            total_rows = old.execute('SELECT COUNT(*) FROM raw_excel_rows').fetchone()[0]
            emit({'stage': '기존 원천 행 비교 준비', 'completed': 0, 'total': total_rows, 'unit': '행'})
            processed = 0
            for row in old.execute('SELECT ' + ','.join(FIELDS) + ' FROM raw_excel_rows'):
                values = dict(row)
                conn.execute('INSERT INTO previous VALUES (?,?)',
                             (values['unit_code'], hashlib.sha256(_payload(values).encode()).hexdigest()))
                processed += 1
                if processed % 4096 == 0:
                    emit({'stage': '기존 원천 행 비교 준비', 'completed': processed, 'total': total_rows, 'unit': '행'})
            emit({'stage': '기존 원천 행 비교 준비', 'completed': processed, 'total': total_rows, 'unit': '행'})
        workbook = load_workbook(excel, read_only=True, data_only=True)
        majors: set[str] = set()
        rows_count = 0
        try:
            total_sheets = len(workbook.worksheets)
            emit({'stage': 'Excel 시트 읽기', 'completed': 0, 'total': total_sheets, 'unit': '시트'})
            for sheet_number, sheet in enumerate(workbook.worksheets, 1):
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    emit({'stage': 'Excel 시트 읽기', 'completed': sheet_number, 'total': total_sheets, 'unit': '시트'})
                    continue
                mapping = build_header_map(header)
                for number, row in enumerate(rows, 2):
                    if not any(value is not None and str(value).strip() for value in row):
                        continue
                    values = {field: get(row, mapping, field) for field in FIELDS}
                    for required in ('major_code', 'unit_code', 'element_code', 'criteria_no'):
                        if not values[required]:
                            raise ValueError(f'{sheet.title}:{number}: missing {required}')
                    payload = _payload(values)
                    conn.execute('INSERT INTO incoming VALUES (?,?,?,?,?)',
                                 (values['unit_code'], hashlib.sha256(payload.encode()).hexdigest(),
                                  payload, sheet.title, number))
                    majors.add(values['major_code'])
                    rows_count += 1
                emit({'stage': 'Excel 시트 읽기', 'completed': sheet_number, 'total': total_sheets, 'unit': '시트'})
        finally:
            workbook.close()
        if not rows_count:
            raise ValueError('Workbook has no NCS data rows; refusing empty replacement')
        emit({'stage': '능력단위 변경분 비교', 'completed': 0, 'total': None, 'unit': '능력단위'})
        conn.executescript('CREATE INDEX incoming_unit ON incoming(unit); CREATE INDEX previous_unit ON previous(unit);')
        previous, incoming = _fingerprints(conn, 'previous'), _fingerprints(conn, 'incoming')
        added = sorted(incoming.keys() - previous.keys())
        removed = sorted(previous.keys() - incoming.keys())
        changed = sorted(unit for unit in incoming.keys() & previous.keys() if incoming[unit] != previous[unit])
        unchanged = sorted(unit for unit in incoming.keys() & previous.keys() if incoming[unit] == previous[unit])
        conn.commit()
        return {'inserted_units': added, 'updated_units': changed, 'deleted_units': removed,
                'unchanged_units': unchanged, 'incoming_rows': rows_count, 'major_codes': sorted(majors),
                'absent_baseline_major_codes': sorted(previous_majors - majors),
                'counts': {'inserted': len(added), 'updated': len(changed), 'deleted': len(removed),
                           'unchanged': len(unchanged)}}, conn
    except BaseException:
        conn.close()
        raise


def _retire(conn: sqlite3.Connection, affected: list[str], removed: list[str],
            *, authorize: Callable[[], None] = lambda: None) -> dict[str, int]:
    """Follow actual foreign keys rather than maintaining an incomplete table list."""
    tables = [row[0] for row in conn.execute("SELECT name, sql FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' AND name NOT LIKE 'builder_%'") if not is_dbstat_table(row[1])]
    conn.execute('CREATE TEMP TABLE delta_units(unit TEXT PRIMARY KEY)')
    conn.executemany('INSERT INTO delta_units VALUES (?)', ((unit,) for unit in affected))
    for table in tables:
        conn.execute(f'CREATE TEMP TABLE {_q("retire_" + table)}(id INTEGER PRIMARY KEY)')
    conn.execute('CREATE TEMP TABLE removed_units(unit TEXT PRIMARY KEY)')
    conn.executemany('INSERT INTO removed_units VALUES (?)', ((unit,) for unit in removed))
    conn.execute('INSERT INTO retire_competency_units SELECT rowid FROM competency_units WHERE unit_code IN (SELECT unit FROM removed_units)')
    for table in ('competency_elements', 'raw_excel_rows'):
        conn.execute(f'INSERT INTO {_q("retire_" + table)} SELECT rowid FROM {_q(table)} WHERE unit_code IN (SELECT unit FROM delta_units)')
    trusted_units = conn.execute(f'SELECT count(*) FROM competency_units WHERE unit_code IN (SELECT unit FROM delta_units) AND review_status IN ({TRUSTED})').fetchone()[0]
    if trusted_units:
        raise ValueError(f'Changed source touches human decisions: competency_units={trusted_units}')
    # API and career source rows remain available when an Excel unit disappears.
    # Only their optional resolved join is detached, never their source payload.
    for table in ('ncs_unit_standard_training', 'ncs_career_paths'):
        if table not in tables:
            continue
        columns = {row[1] for row in conn.execute(f'PRAGMA table_info({_q(table)})')}
        if 'matched_unit_code' not in columns:
            continue
        if 'review_status' in columns and conn.execute(f'SELECT count(*) FROM {_q(table)} WHERE matched_unit_code IN (SELECT unit FROM removed_units) AND review_status IN ({TRUSTED})').fetchone()[0]:
            raise ValueError(f'Changed source touches human decisions: {table}')
        conn.execute(f'UPDATE {_q(table)} SET matched_unit_code=NULL WHERE matched_unit_code IN (SELECT unit FROM removed_units)')
    # Derived summaries do not reference the source relations they summarize.
    # Reconcile automatic rows globally; preserve unrelated human decisions.
    aggregate_tables = ('task_similarity_links', 'ncs_training_course_concept_links',
                        'ncs_training_course_element_links', 'training_goal_concept_links',
                        'training_delivery_relations', 'ncs_training_course_unit_links')
    for table in aggregate_tables:
        if table in tables:
            conn.execute(f'INSERT OR IGNORE INTO {_q("retire_" + table)} SELECT rowid FROM {_q(table)} WHERE review_status NOT IN ({TRUSTED})')
    if 'ontology_concept_relations' in tables:
        conn.execute(f"INSERT OR IGNORE INTO retire_ontology_concept_relations SELECT rowid FROM ontology_concept_relations WHERE review_status NOT IN ({TRUSTED}) AND relation_type IN ('co_required_in_element','knowledge_enables_skill','attitude_supports_skill','knowledge_informs_attitude')")
    foreign_keys = [(table, fk) for table in tables for fk in conn.execute(f'PRAGMA foreign_key_list({_q(table)})') if fk[2] in tables]
    while True:
        before = conn.total_changes
        for table, fk in foreign_keys:
            parent, child_col, parent_col = fk[2], fk[3], fk[4]
            if not parent_col:
                parent_col = next(row[1] for row in conn.execute(f'PRAGMA table_info({_q(parent)})') if row[5])
            conn.execute(f'INSERT OR IGNORE INTO {_q("retire_" + table)} SELECT c.rowid FROM {_q(table)} c JOIN {_q(parent)} p ON c.{_q(child_col)}=p.{_q(parent_col)} JOIN {_q("retire_" + parent)} r ON r.id=p.rowid')
        if conn.total_changes == before:
            break
    blocked = {}
    for table in tables:
        columns = {row[1] for row in conn.execute(f'PRAGMA table_info({_q(table)})')}
        status_columns = columns.intersection(('review_status', 'link_status', 'status'))
        if status_columns:
            status_clause = ' OR '.join(f'{_q(column)} IN ({TRUSTED})' for column in sorted(status_columns))
            count = conn.execute(f'SELECT count(*) FROM {_q(table)} WHERE rowid IN (SELECT id FROM {_q("retire_" + table)}) AND ({status_clause})').fetchone()[0]
            if count:
                blocked[table] = count
    if blocked:
        raise ValueError('Changed source touches human decisions; explicit review required: ' + json.dumps(blocked))
    conn.execute('CREATE TABLE IF NOT EXISTS builder_retired_rows(archive_id INTEGER PRIMARY KEY, retired_at TEXT NOT NULL, table_name TEXT NOT NULL, original_rowid INTEGER NOT NULL, row_json TEXT NOT NULL)')
    retired = {}
    for table in tables:
        count = conn.execute(f'SELECT count(*) FROM {_q("retire_" + table)}').fetchone()[0]
        if not count:
            continue
        retired[table] = count
        # Derived automatic edges can be regenerated; archiving all global
        # summaries would needlessly double multi-GB graph storage.
        columns = {row[1] for row in conn.execute(f'PRAGMA table_info({_q(table)})')}
        if table not in ('raw_excel_rows', 'competency_units', 'competency_elements',
                         'performance_criteria', 'ksa_items') and 'source_payload' not in columns:
            continue
        rows = conn.execute(f'SELECT rowid AS __rowid__, * FROM {_q(table)} WHERE rowid IN (SELECT id FROM {_q("retire_" + table)})')
        for row in rows:
            item = dict(row)
            original = item.pop('__rowid__')
            conn.execute('INSERT INTO builder_retired_rows(retired_at,table_name,original_rowid,row_json) VALUES (?,?,?,?)',
                         (now_utc(), table, original, json.dumps(item, ensure_ascii=False, default=str)))
    authorize()
    conn.commit()
    conn.execute('PRAGMA foreign_keys=OFF')
    for table in retired:
        authorize()
        conn.execute(f'DELETE FROM {_q(table)} WHERE rowid IN (SELECT id FROM {_q("retire_" + table)})')
    authorize()
    conn.commit()
    conn.execute('PRAGMA foreign_keys=ON')
    return retired


def build_excel_delta(excel_path: Path, baseline_db: Path, output_db: Path,
                      work_dir: Path, progress: Callable[[str | dict[str, Any]], None] | None = None,
                      *, builder_context: BuilderOperationContext) -> dict[str, Any]:
    """Treat the uploaded workbook as a complete replacement source snapshot.

    Source normalization is unit-incremental. Existing ontology algorithms skip
    existing atomic KSA but recompute cross-unit/task/training aggregates globally.
    A candidate is published at output_db only after foreign-key/integrity checks.
    """
    output_path = Path(os.path.abspath(output_db))
    requested_work_dir = Path(work_dir)
    work_path = Path(os.path.abspath(requested_work_dir))

    def authorize() -> None:
        context = require_builder_context(
            builder_context, action="build_delta", version_dir=output_path.parent,
        )
        expected_work_dir = Path(context.state_dir) / 'versions' / context.version / 'delta'
        if ('..' in requested_work_dir.parts
                or work_path != output_path.parent / 'delta'
                or work_path.resolve() != expected_work_dir):
            raise BuilderAuthorizationError('Excel work directory must be the exact Builder version delta directory.')
        try:
            info = work_path.lstat()
        except FileNotFoundError:
            pass
        else:
            if (not stat.S_ISDIR(info.st_mode) or stat.S_ISLNK(info.st_mode)
                    or getattr(info, 'st_file_attributes', 0) & 0x400):
                raise BuilderAuthorizationError('Excel work directory must not be a link or reparse point.')
        try:
            output_path.lstat()
        except FileNotFoundError:
            pass
        else:
            raise ValueError('Candidate output must be a new path separate from input files')

    authorize()
    excel, baseline, output = (Path(path).resolve() for path in (excel_path, baseline_db, output_path))
    if not excel.is_file() or not baseline.is_file():
        raise FileNotFoundError('Excel and baseline database must exist')
    if output.exists() or output in (excel, baseline):
        raise ValueError('Candidate output must be a new path separate from input files')
    emit = progress or (lambda message: None)
    authorize()
    work_path.mkdir(parents=True, exist_ok=True)
    authorize()
    output.parent.mkdir(parents=True, exist_ok=True)
    authorize()
    with tempfile.TemporaryDirectory(prefix='excel-delta-', dir=work_path) as temporary:
        candidate = Path(temporary) / 'candidate.db'
        emit('Creating a consistent read-only baseline snapshot')
        _sqlite_online_snapshot(baseline, candidate, progress=progress, authorize=authorize)
        emit('Comparing workbook content by competency unit')
        authorize()
        plan, stage = _stage(excel, candidate, Path(temporary) / 'rows.db', progress=progress)
        try:
            affected = sorted(set(plan['inserted_units'] + plan['updated_units'] + plan['deleted_units']))
            stages, retired = [], {}
            if affected:
                emit(f'Normalizing {len(affected)} changed units in a separate candidate')
                authorize()
                conn = connect(candidate)
                try:
                    initialize_database(conn)
                    emit({'stage': '변경 원천 의존 관계 정리', 'completed': 0, 'total': None, 'unit': '단계'})
                    authorize()
                    retired = _retire(conn, affected, plan['deleted_units'], authorize=authorize)
                    normalizer = Normalizer(conn, excel.name)
                    timestamp = now_utc()
                    units_to_normalize = plan['inserted_units'] + plan['updated_units']
                    emit({'stage': '변경 능력단위 정규화', 'completed': 0, 'total': len(units_to_normalize), 'unit': '능력단위'})
                    for unit_number, unit in enumerate(units_to_normalize, 1):
                        authorize()
                        for payload, sheet, number in stage.execute('SELECT payload,sheet,row_number FROM incoming WHERE unit=?', (unit,)):
                            values = dict(zip(FIELDS, json.loads(payload)))
                            classification_id = normalizer.get_classification_id(values)
                            existing_classification = conn.execute('SELECT * FROM classifications WHERE classification_id=?', (classification_id,)).fetchone()
                            name_fields = ('major_name', 'middle_name', 'small_name', 'sub_name')
                            if existing_classification['review_status'] in ('human_reviewed', 'accepted', 'reviewed') and any(existing_classification[key] != values[key] for key in name_fields):
                                raise ValueError(f'Changed classification {classification_id} touches human decisions')
                            conn.execute('UPDATE classifications SET major_name=?,middle_name=?,small_name=?,sub_name=? WHERE classification_id=?',
                                         (*(values[key] for key in ('major_name','middle_name','small_name','sub_name')), classification_id))
                            normalizer.ingest(sheet, number, values, timestamp)
                        emit({'stage': '변경 능력단위 정규화', 'completed': unit_number, 'total': len(units_to_normalize), 'unit': '능력단위'})
                    authorize()
                    conn.commit()
                finally:
                    conn.close()
                emit('Rebuilding missing ontology nodes and reconciling dependent global relations')
                stages, invariants = _run_pipeline(candidate, bootstrap=True, progress=progress, authorize=authorize)
            else:
                invariants = {'raw_ksa_preserved': True, 'trusted_statuses_preserved': True}
                emit('No content changes; ontology preprocessing skipped')
                emit({'stage': '원천 변경 없음 · 온톨로지 전처리 생략', 'completed': len(plan['unchanged_units']),
                      'total': len(plan['unchanged_units']), 'unit': '능력단위'})
            with closing(sqlite3.connect(candidate)) as check:
                emit({'stage': '후보 DB 무결성·참조 검증', 'completed': 0, 'total': None, 'unit': '검사'})
                integrity = check.execute('PRAGMA integrity_check').fetchone()[0]
                violations = check.execute('PRAGMA foreign_key_check').fetchmany(20)
                if integrity != 'ok' or violations:
                    raise ValueError(f'Candidate validation failed: {integrity}; foreign keys: {violations}')
                authorize()
                check.execute('PRAGMA wal_checkpoint(TRUNCATE)')
                authorize()
                check.execute('PRAGMA journal_mode=DELETE')
            candidate_integrity = {
                'sha256': file_sha256(candidate),
                'raw_ksa_sha256': raw_ksa_sha256(candidate),
                'trusted_review_status_identity_digest': trusted_review_status_identity_digest(candidate),
            }
            authorize()
            os.replace(candidate, output)
            return {'schema': 'ncs_excel_delta_build_v1', 'ok': True, 'success': True,
                    'status': 'candidate_ready', 'candidate_db': str(output), 'output_db': str(output),
                    'baseline_db': str(baseline), 'source_delta': plan, 'affected_units': affected,
                    'retired_rows': retired, 'stages': stages, 'invariants': invariants,
                    'candidate_integrity': candidate_integrity,
                    'source_normalization': 'changed_units_only',
                    'ontology_processing': 'incremental_nodes_global_dependent_relations' if affected else 'skipped_no_change',
                    'source_snapshot_semantics': 'complete_workbook_replacement',
                    'baseline_modified': False, 'approval_claim': False,
                    'validation': {'integrity_check': integrity, 'foreign_key_violations': 0}}
        finally:
            stage.close()

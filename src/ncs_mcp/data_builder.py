"""Versioned local Excel/API builds. Serving databases are never rebuilt in place."""

from __future__ import annotations

import json
import os
import shutil
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

from .api_refresh_builder import (
    file_sha256,
    refresh_ncs_api_evidence,
    raw_ksa_sha256,
    trusted_review_status_identity_digest,
    trusted_review_status_counts,
)
from .config import PROJECT_ROOT, load_settings
from .builder_authorization import (
    BuilderAuthorizationError,
    BuilderOperationContext,
    _bind_operation_version,
    _exclusive_operation,
    require_builder_context,
    validate_builder_paths,
)
from .ontology_refresh_builder import _sqlite_online_snapshot
from .preprocess_excel import build_header_map


class BuilderError(RuntimeError):
    """User-facing operational failure without credentials or API response bodies."""


def _absolute_path(path: str | Path) -> Path:
    """Make *path* absolute without expanding Windows 8.3 path spelling."""

    return Path(os.path.abspath(Path(path).expanduser()))


def api_failure_message(evidence: dict) -> str:
    phase_names = {
        "source_invariant_check": "원본 DB 검사",
        "working_copy_backup": "작업 DB 복사",
        "working_copy_invariant_check": "작업 DB 검사",
        "api_collection": "API 수집",
        "post_collection_invariant_check": "수집 후 데이터 검사",
        "training_link_build": "교육과정 연결",
        "final_invariant_check": "최종 원본 보존 검사",
    }
    reason_names = {
        "sqlite_dbstat_module_unavailable": "SQLite 진단 모듈(dbstat) 호환성 오류",
        "sqlite_database_locked": "DB가 다른 작업에서 사용 중입니다",
        "disk_full": "디스크 여유 공간이 부족합니다",
        "sqlite_read_only": "작업 경로에 쓰기 권한이 없습니다",
        "local_database_or_filesystem_error": "로컬 DB 또는 파일 처리 오류",
    }
    details = list(evidence.get("preflight_errors") or [])
    if evidence.get("failed_phase"):
        details.append(
            phase_names.get(evidence["failed_phase"], evidence["failed_phase"])
        )
    if evidence.get("failure_reason"):
        details.append(
            reason_names.get(evidence["failure_reason"], evidence["failure_reason"])
        )
    if evidence.get("failure_type"):
        details.append(evidence["failure_type"])
    for source, rows in (evidence.get("source_results") or {}).items():
        failed = [row for row in rows if not row.get("completion_proven")]
        if failed:
            examples = ", ".join(
                f"{row.get('major_code')}:{row.get('error_type', '응답 완결성 미확인')}"
                for row in failed[:3]
            )
            details.append(f"{source} 실패/미확인 {len(failed)}개 대분류 ({examples})")
    link = evidence.get("training_link_build") or {}
    if link.get("error_type"):
        details.append("교육과정 연결: " + link["error_type"])
    for key in ("invariant_failure", "source_integrity_failure"):
        if evidence.get(key):
            details.append("원문 또는 검토 상태 보존 검사 실패")
    return "API 갱신 중단: " + " / ".join(
        details or [evidence.get("outcome", "원인 미기록")]
    )


def atomic_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    try:
        temporary.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def inspect_workbook(path: Path) -> dict:
    path = Path(path).resolve(strict=True)
    if path.suffix.lower() != ".xlsx":
        raise BuilderError("NCS 정보망 .xlsx 파일을 선택하세요.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            try:
                build_header_map(header)
            except ValueError as exc:
                raise BuilderError(
                    f"시트 '{sheet.title}'의 필수 NCS 열이 맞지 않습니다: {exc}"
                ) from exc
            first = next(rows, None)
            sheets.append(
                {
                    "name": sheet.title,
                    "estimated_rows": max(0, (sheet.max_row or 1) - 1),
                    "has_data": first is not None and any(v is not None for v in first),
                }
            )
    finally:
        workbook.close()
    if not sheets or not any(s["has_data"] for s in sheets):
        raise BuilderError("처리할 NCS 데이터 행이 없습니다.")
    return {
        "filename": path.name,
        "bytes": path.stat().st_size,
        "sheets": sheets,
        "note": "행 수는 Excel 메타데이터 추정치입니다. 실제 건수는 전처리 후 확정됩니다.",
    }


class DataBuilder:
    def __init__(
        self, root: Path = PROJECT_ROOT, progress: Callable[[str], None] | None = None
    ):
        # Keep the caller's absolute path spelling.  On GitHub Windows runners,
        # ``resolve()`` expands RUNNER~1 to runneradmin, which makes paths to the
        # same file compare differently across Builder reports and call sites.
        # Individual containment checks still use resolved paths.
        self.root = _absolute_path(root)
        self.state = self.root / ".state/ncs-data-builder"
        validate_builder_paths(self.root, self.state)
        self.state.mkdir(parents=True, exist_ok=True)
        validate_builder_paths(self.root, self.state)
        self.progress = progress or (lambda message: None)

    @contextmanager
    def exclusive(self, action: str = "exclusive", version: str | None = None):
        try:
            with _exclusive_operation(
                root=self.root, state_dir=self.state, action=action, version=version
            ) as context:
                yield context
        except BuilderAuthorizationError as exc:
            raise BuilderError(
                "Builder 작업 권한을 확인하지 못했습니다. 다른 작업이 실행 중이거나 "
                "operation.lock이 변경되었습니다. 비정상 종료했다면 실행 프로세스를 먼저 확인하세요."
            ) from exc

    def _version_dir(self, version: str) -> Path:
        if not version or any(c not in "0123456789abcdef_-" for c in version):
            raise BuilderError("올바르지 않은 버전 ID입니다.")
        directory = self.state / "versions" / version
        validate_builder_paths(self.root, self.state, directory)
        directory = directory.resolve()
        return directory

    def versions(self) -> list[dict]:
        result = []
        for path in sorted(
            (self.state / "versions").glob("*/build.json"), reverse=True
        ):
            try:
                result.append(json.loads(path.read_text(encoding="utf-8")))
            except (ValueError, OSError):
                continue
        return result

    def _new(
        self, kind: str, builder_context: BuilderOperationContext
    ) -> tuple[Path, dict]:
        require_builder_context(
            builder_context,
            action={"excel-delta": "build_delta", "api": "refresh_api", "current-copy": "copy_current"}[kind],
            root=self.root, state_dir=self.state,
        )
        version = (
            datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
            + "_"
            + uuid.uuid4().hex[:8]
        )
        folder = self._version_dir(version)
        _bind_operation_version(builder_context, version)
        require_builder_context(builder_context, action=builder_context.action, version_dir=folder)
        folder.mkdir(parents=True)
        report = {
            "schema": "ncs_data_builder_version_v1",
            "version": version,
            "kind": kind,
            "status": "building",
            "created_at": datetime.now(timezone.utc).isoformat(),
            "human_approval_claim": False,
            "operation_lineage": builder_context.lineage(),
        }
        self._write_version_json(folder, "build.json", report, builder_context,
                                 action=builder_context.action, version=version)
        return folder, report

    def _write_version_json(
        self, folder: Path, filename: str, payload: dict,
        builder_context: BuilderOperationContext, *, action: str | tuple[str, ...], version: str,
    ) -> None:
        """Write Builder reports only while their exact version lease is live."""
        if filename not in {"build.json", "api-refresh.json"}:
            raise BuilderAuthorizationError("Unsupported Builder version report.")

        def authorize() -> None:
            require_builder_context(
                builder_context, action=action, root=self.root, state_dir=self.state,
                version=version, version_dir=folder,
            )

        authorize()
        path = folder / filename
        temporary = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
        text = json.dumps(payload, ensure_ascii=False, indent=2)
        try:
            authorize()
            with temporary.open("x", encoding="utf-8") as stream:
                stream.write(text)
                stream.flush()
                os.fsync(stream.fileno())
            authorize()
            os.replace(temporary, path)
        finally:
            # Losing a lease can also mean losing the folder. Preserve a pending
            # temp as evidence instead of mutating a replacement owner's path.
            try:
                authorize()
            except BuilderAuthorizationError:
                pass
            else:
                temporary.unlink(missing_ok=True)

    def _space(self, required: int) -> None:
        if shutil.disk_usage(self.state).free < required:
            raise BuilderError(
                f"작업 공간이 부족합니다. 최소 {required / 1024**3:.1f} GB 여유 공간이 필요합니다."
            )

    def _finish(self, folder: Path, report: dict,
                builder_context: BuilderOperationContext) -> dict:
        from .builder_validation import validate_candidate

        def authorize():
            require_builder_context(
                builder_context, action=("build_delta", "refresh_api", "copy_current", "resume"),
                root=self.root, state_dir=self.state, version=report["version"], version_dir=folder,
            )

        authorize()
        report["checkpoint"] = "final_validation"
        self._write_version_json(folder, "build.json", report, builder_context,
                                 action=builder_context.action, version=report["version"])
        try:
            authorize()
            result = validate_candidate(
                folder / "ncs.db", folder / "validation-checkpoint.json", self.progress,
                builder_context=builder_context,
            )
        except ValueError as exc:
            raise BuilderError(str(exc)) from exc
        authorize()
        report.update(status="ready", **result)
        self._write_version_json(folder, "build.json", report, builder_context,
                                 action=builder_context.action, version=report["version"])
        return report

    def _failed(self, folder: Path, report: dict, exc: Exception,
                builder_context: BuilderOperationContext) -> None:
        # Do not serialize arbitrary exceptions: requests exceptions can contain service keys.
        try:
            require_builder_context(builder_context,
                                    action=("build_delta", "refresh_api", "resume", "copy_current"),
                                    root=self.root, state_dir=self.state,
                                    version=report["version"], version_dir=folder)
            report.update(status="failed", error_type=type(exc).__name__)
            self._write_version_json(folder, "build.json", report, builder_context,
                                     action=builder_context.action, version=report["version"])
        except BuilderAuthorizationError:
            return

    def candidate(self, version: str) -> Path:
        folder = self._version_dir(version)
        report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
        if report.get("status") != "ready":
            raise BuilderError("검증을 통과한 버전만 사용할 수 있습니다.")
        db = folder / "ncs.db"
        if file_sha256(db, self.progress, "선택 버전 DB 검사") != report.get("sha256"):
            raise BuilderError("검증 이후 DB가 변경되었습니다. 다시 빌드하세요.")
        return db

    def gold_preflight(self, version: str) -> dict:
        """Inspect Gold capacity for a verified version without exporting it."""

        from .builder_gold import gold_preflight

        db = self.candidate(version)
        report = json.loads(
            (self._version_dir(version) / "build.json").read_text(encoding="utf-8")
        )
        return gold_preflight(
            builder_version=version,
            db_path=db,
            source_db_sha256=str(report["sha256"]),
        )

    def prepare_gold(
        self,
        version: str,
        *,
        previous_version: str | None = None,
        batch_size: int = 10_000,
        internal_roles=(),
        role_alignments=(),
        role_mapping_packet: str | Path | None = None,
    ) -> dict:
        """Build a validated streaming Gold artifact for an existing version."""

        from .builder_gold import (
            GOLD_EXPORT_NAME,
            load_internal_role_mapping_packet,
            prepare_builder_gold,
        )

        with self.exclusive("prepare_gold", version) as builder_context:
            db = self.candidate(version)
            folder = self._version_dir(version)
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            previous_export = None
            if previous_version is not None:
                self.candidate(previous_version)
                previous_export = (
                    self._version_dir(previous_version) / "gold" / GOLD_EXPORT_NAME
                )
                if not previous_export.is_file():
                    raise BuilderError(
                        "Previous Builder version has no prepared Gold export."
                    )
            try:
                if role_mapping_packet is not None:
                    if tuple(internal_roles) or tuple(role_alignments):
                        raise ValueError(
                            "role_mapping_packet cannot be combined with explicit role inputs"
                        )
                    internal_roles, role_alignments = load_internal_role_mapping_packet(
                        role_mapping_packet
                    )
                return prepare_builder_gold(
                    builder_version=version,
                    version_dir=folder,
                    db_path=db,
                    source_db_sha256=str(report["sha256"]),
                    batch_size=batch_size,
                    previous_export=previous_export,
                    internal_roles=internal_roles,
                    role_alignments=role_alignments,
                )
            except (OSError, ValueError, RuntimeError) as exc:
                raise BuilderError(
                    f"Gold preparation failed ({type(exc).__name__})."
                ) from exc

    def sync_gold(
        self,
        version: str,
        *,
        apply: bool = False,
        reconcile: bool = False,
        batch_size: int = 1_000,
        settings=None,
    ) -> dict:
        """Dry-run validate or explicitly apply a prepared Gold artifact."""

        from .builder_gold import sync_builder_gold

        with self.exclusive("sync_gold", version) as builder_context:
            self.candidate(version)
            folder = self._version_dir(version)
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            try:
                return sync_builder_gold(
                    builder_version=version,
                    version_dir=folder,
                    source_db_sha256=str(report["sha256"]),
                    apply=apply,
                    reconcile=reconcile,
                    batch_size=batch_size,
                    settings=settings,
                )
            except (OSError, ValueError, RuntimeError) as exc:
                raise BuilderError(f"Gold sync failed ({type(exc).__name__}).") from exc

    def prepare_gold_embeddings(
        self,
        version: str,
        *,
        model: str | None = None,
        dimensions: int | None = None,
        device: str | None = None,
        allow_download: bool = False,
        entity_types=None,
        batch_size: int = 64,
        fetch_size: int = 128,
        max_text_chars: int = 8_192,
        max_records: int | None = 20,
        provider=None,
    ) -> dict:
        """Create a local, source-bound Gold embedding patch for one version."""

        from .builder_gold import (
            DEFAULT_GOLD_EMBEDDING_DIMENSIONS,
            DEFAULT_GOLD_EMBEDDING_MODEL,
            prepare_builder_gold_embeddings,
        )
        from .embedding_batches import SUPPORTED_ENTITY_TYPES
        from .local_embeddings import SentenceTransformerEmbeddingProvider

        with self.exclusive("prepare_gold_embeddings", version) as builder_context:
            db = self.candidate(version)
            folder = self._version_dir(version)
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            try:
                if provider is None:
                    configured_model = (
                        str(model or "").strip()
                        or os.environ.get(
                            "NCS_MCP_GOLD_LOCAL_EMBEDDING_MODEL", ""
                        ).strip()
                        or DEFAULT_GOLD_EMBEDDING_MODEL
                    )
                    configured_dimensions = dimensions
                    if configured_dimensions is None:
                        raw_dimensions = os.environ.get(
                            "NCS_MCP_GOLD_EMBEDDING_DIMENSIONS", ""
                        ).strip()
                        configured_dimensions = (
                            int(raw_dimensions)
                            if raw_dimensions
                            else DEFAULT_GOLD_EMBEDDING_DIMENSIONS
                        )
                    configured_device = (
                        str(device or "").strip()
                        or os.environ.get(
                            "NCS_MCP_GOLD_LOCAL_EMBEDDING_DEVICE", ""
                        ).strip()
                        or None
                    )
                    provider = SentenceTransformerEmbeddingProvider(
                        configured_model,
                        dimensions=configured_dimensions,
                        device=configured_device,
                        local_files_only=not allow_download,
                    )
                return prepare_builder_gold_embeddings(
                    builder_version=version,
                    version_dir=folder,
                    db_path=db,
                    source_db_sha256=str(report["sha256"]),
                    provider=provider,
                    entity_types=tuple(entity_types or SUPPORTED_ENTITY_TYPES),
                    batch_size=batch_size,
                    fetch_size=fetch_size,
                    max_text_chars=max_text_chars,
                    max_records=max_records,
                )
            except (OSError, ValueError, RuntimeError, TypeError) as exc:
                raise BuilderError(
                    f"Gold embedding preparation failed ({type(exc).__name__})."
                ) from exc

    def sync_gold_embeddings(
        self,
        version: str,
        *,
        apply: bool = False,
        create_indexes: bool = True,
        batch_size: int = 1_000,
        max_retries: int = 2,
        settings=None,
    ) -> dict:
        """Dry-run validate or explicitly apply one Builder embedding patch."""

        from .builder_gold import sync_builder_gold_embeddings

        with self.exclusive("sync_gold_embeddings", version) as builder_context:
            self.candidate(version)
            folder = self._version_dir(version)
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            try:
                return sync_builder_gold_embeddings(
                    builder_version=version,
                    version_dir=folder,
                    source_db_sha256=str(report["sha256"]),
                    apply=apply,
                    create_indexes=create_indexes,
                    batch_size=batch_size,
                    max_retries=max_retries,
                    settings=settings,
                )
            except (OSError, ValueError, RuntimeError, TypeError) as exc:
                raise BuilderError(
                    f"Gold embedding sync failed ({type(exc).__name__})."
                ) from exc

    def prepare_gold_embedding_shards(
        self,
        version: str,
        *,
        model: str | None = None,
        dimensions: int | None = None,
        device: str | None = None,
        allow_download: bool = False,
        provider=None,
        entity_types=None,
        shard_size: int = 10_000,
        batch_size: int = 64,
        fetch_size: int = 128,
        max_text_chars: int = 8_192,
    ) -> dict:
        """Prepare resumable full shards; unlike the smoke API this has no cap."""

        from .builder_gold import (
            DEFAULT_GOLD_EMBEDDING_DIMENSIONS,
            DEFAULT_GOLD_EMBEDDING_MODEL,
            prepare_builder_gold_embedding_shards,
        )
        from .embedding_batches import SUPPORTED_ENTITY_TYPES
        from .local_embeddings import SentenceTransformerEmbeddingProvider

        with self.exclusive("prepare_gold_embedding_shards", version) as builder_context:
            db = self.candidate(version)
            folder = self._version_dir(version)
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            try:
                if provider is None:
                    configured_model = (
                        str(model or "").strip()
                        or os.environ.get(
                            "NCS_MCP_GOLD_LOCAL_EMBEDDING_MODEL", ""
                        ).strip()
                        or DEFAULT_GOLD_EMBEDDING_MODEL
                    )
                    configured_dimensions = dimensions
                    if configured_dimensions is None:
                        raw_dimensions = os.environ.get(
                            "NCS_MCP_GOLD_EMBEDDING_DIMENSIONS", ""
                        ).strip()
                        configured_dimensions = (
                            int(raw_dimensions)
                            if raw_dimensions
                            else DEFAULT_GOLD_EMBEDDING_DIMENSIONS
                        )
                    configured_device = (
                        str(device or "").strip()
                        or os.environ.get(
                            "NCS_MCP_GOLD_LOCAL_EMBEDDING_DEVICE", ""
                        ).strip()
                        or None
                    )
                    provider = SentenceTransformerEmbeddingProvider(
                        configured_model,
                        dimensions=configured_dimensions,
                        device=configured_device,
                        local_files_only=not allow_download,
                    )
                return prepare_builder_gold_embedding_shards(
                    builder_version=version,
                    version_dir=folder,
                    db_path=db,
                    source_db_sha256=str(report["sha256"]),
                    provider=provider,
                    entity_types=tuple(entity_types or SUPPORTED_ENTITY_TYPES),
                    shard_size=shard_size,
                    batch_size=batch_size,
                    fetch_size=fetch_size,
                    max_text_chars=max_text_chars,
                )
            except (OSError, ValueError, RuntimeError, TypeError) as exc:
                raise BuilderError(
                    f"Gold embedding shard preparation failed ({type(exc).__name__})."
                ) from exc

    def sync_gold_embedding_shards(
        self,
        version: str,
        *,
        apply: bool = False,
        create_indexes: bool = True,
        batch_size: int = 1_000,
        max_retries: int = 2,
        settings=None,
    ) -> dict:
        """Dry-run or explicitly apply a prepared resumable shard manifest."""

        from .builder_gold import sync_builder_gold_embedding_shards

        with self.exclusive("sync_gold_embedding_shards", version) as builder_context:
            self.candidate(version)
            folder = self._version_dir(version)
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            try:
                return sync_builder_gold_embedding_shards(
                    builder_version=version,
                    version_dir=folder,
                    source_db_sha256=str(report["sha256"]),
                    apply=apply,
                    create_indexes=create_indexes,
                    batch_size=batch_size,
                    max_retries=max_retries,
                    settings=settings,
                )
            except (OSError, ValueError, RuntimeError, TypeError) as exc:
                raise BuilderError(
                    f"Gold embedding shard sync failed ({type(exc).__name__})."
                ) from exc

    def build_delta(self, source: Path, baseline: Path) -> dict:
        from .excel_delta_builder import build_excel_delta

        with self.exclusive("build_delta") as builder_context:
            preview = inspect_workbook(source)
            baseline = Path(baseline).resolve(strict=True)
            self._space(
                baseline.stat().st_size * 2 + max(1024**3, preview["bytes"] * 10)
            )
            folder, report = self._new("excel-delta", builder_context)
            try:
                original = folder / "source.xlsx"
                self.progress("새 Excel 원본 보관 및 능력단위 변경분 비교")
                require_builder_context(builder_context, action="build_delta", version_dir=folder)
                shutil.copyfile(source, original)
                report.update(
                    source=preview,
                    source_sha256=file_sha256(original),
                    parent_database=str(baseline),
                )
                delta = build_excel_delta(
                    original,
                    baseline,
                    folder / "ncs.db",
                    folder / "delta",
                    self.progress,
                    builder_context=builder_context,
                )
                require_builder_context(builder_context, action="build_delta", version_dir=folder)
                atomic_json(folder / "delta.json", delta)
                if not delta.get("ok"):
                    raise BuilderError(
                        "변경분 검증에 실패했습니다. delta.json을 확인하세요."
                    )
                report["source_delta"] = delta.get("source_delta")
                report["ontology_processing"] = delta.get("ontology_processing")
                return self._finish(folder, report, builder_context)
            except Exception as exc:
                self._failed(folder, report, exc, builder_context)
                if isinstance(exc, ValueError):
                    raise BuilderError(str(exc)) from exc
                raise

    def package(self, version: str, deploy_root: Path) -> dict:
        from .builder_release import build_release

        with self.exclusive("package", version) as builder_context:
            db = self.candidate(version)
            report = build_release(
                self._version_dir(version),
                repo_root=self.root,
                deploy_root=deploy_root,
                expected_source_sha256=file_sha256(db),
                progress=self.progress,
                builder_context=builder_context,
            )
            if not report.get("ok"):
                raise BuilderError(
                    "경량 패키지 생성·검증에 실패했습니다. release.json을 확인하세요."
                )
            return {"version": version, "package": report}

    def deploy(self, version: str, deploy_root: Path, production_url: str) -> dict:
        from .builder_release import ReleaseGuard, _write, deploy_release

        with self.exclusive("deploy", version) as builder_context:
            release_guard = ReleaseGuard(builder_context, action="deploy",
                                         version_dir=self._version_dir(version), repo_root=self.root)
            self.candidate(version)
            folder = self._version_dir(version)
            release_path = folder / "release.json"
            if not release_path.exists():
                raise BuilderError("먼저 선택 프로젝트의 경량 DB 패키지를 생성하세요.")
            release = json.loads(release_path.read_text(encoding="utf-8"))
            # Package carries its exact project. UI project changes require a new package.
            project = json.loads(
                (Path(deploy_root) / ".vercel/project.json").read_text(encoding="utf-8")
            )
            bound_project = (
                release.get("project") or release.get("project_configuration") or {}
            )
            if bound_project.get("projectId") != project.get("projectId"):
                raise BuilderError("패키지를 만든 프로젝트와 선택 프로젝트가 다릅니다.")
            report = deploy_release(
                folder, production_mcp_url=production_url, progress=self.progress,
                builder_context=builder_context,
            )
            if not report.get("ok"):
                if report.get("status") == "reconciliation_failed":
                    raise BuilderError(
                        "완료된 운영 배포의 재검증에 실패했습니다. 기존 완료 기록과 로컬 기준점은 "
                        "보존했습니다. 운영 배포와 패키지 무결성을 확인한 뒤 다시 시도하세요."
                    )
                raise BuilderError(
                    "Vercel 갱신 검증을 완료하지 못했습니다. release.json의 배포 상태를 확인하세요."
                )
            _write(
                self.state / "deployed.json",
                {
                    "version": version,
                    "production_url": production_url,
                    "deploy_root": str(Path(deploy_root).resolve()),
                    "operation_lineage": builder_context.lineage(),
                    "build_id": report.get("build_id"),
                    "source_sha256": report.get("source_sha256"),
                    "deployment_identity": report.get("production_after_promotion"),
                    "release_report_sha256": file_sha256(release_path),
                },
                builder_context=builder_context, guard=release_guard,
            )
            return {"version": version, "deployment": report}

    def refresh_api(self, source: Path, sources: list[str]) -> dict:
        with self.exclusive("refresh_api") as builder_context:
            source = Path(source).resolve(strict=True)
            self._space(source.stat().st_size * 2 + 1024**3)
            folder, report = self._new("api", builder_context)
            try:
                self.progress("API 사전 점검 및 별도 작업 DB 생성")
                report["parent_database"] = str(source)
                report["sources"] = sources
                self._write_version_json(folder, "build.json", report, builder_context,
                                         action="refresh_api", version=report["version"])
                require_builder_context(
                    builder_context, action="refresh_api", root=self.root, state_dir=self.state,
                    version=report["version"], version_dir=folder,
                )
                evidence = refresh_ncs_api_evidence(
                    source,
                    sources=sources,
                    apply=True,
                    output_path=folder / "ncs.db",
                    progress=self.progress,
                    checkpoint_dir=folder / "api-checkpoint",
                    builder_context=builder_context,
                )
                self._write_version_json(folder, "api-refresh.json", evidence, builder_context,
                                         action="refresh_api", version=report["version"])
                if evidence.get("outcome") not in {
                    "succeeded_append_only",
                    "completed_with_warnings",
                }:
                    raise BuilderError(
                        api_failure_message(evidence)
                        + f"\n상세 보고서: {folder / 'api-refresh.json'}"
                    )
                report["api_outcome"] = evidence["outcome"]
                report["sources"] = sources
                parent_report = source.parent / "build.json"
                if (
                    source.is_relative_to(self.state / "versions")
                    and parent_report.exists()
                ):
                    parent = json.loads(parent_report.read_text(encoding="utf-8"))
                    report["parent_version"] = parent.get("version")
                    for field in ("source_delta", "ontology_processing"):
                        if field in parent:
                            report[field] = parent[field]
                return self._finish(folder, report, builder_context)
            except Exception as exc:
                self._failed(folder, report, exc, builder_context)
                raise

    def current_db(self) -> Path:
        pointer = self.state / "deployed.json"
        if pointer.exists():
            payload = json.loads(pointer.read_text(encoding="utf-8"))
            return self.candidate(payload["version"])
        if self.root.resolve() == PROJECT_ROOT.resolve():
            return load_settings().db_path
        return self.root / "data/processed/ncs.db"

    def resume_kind(self, version: str) -> str | None:
        """Read saved evidence and verify Excel digests without declaring readiness."""
        folder = self._version_dir(version)
        try:
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            if report.get("status") == "ready" or not (folder / "ncs.db").is_file():
                return None
            if report.get("kind") == "api":
                evidence_path = folder / "api-refresh.json"
                evidence = (
                    json.loads(evidence_path.read_text(encoding="utf-8"))
                    if evidence_path.exists()
                    else {}
                )
                if not isinstance(evidence, dict):
                    return None
                working_invariants = evidence.get("working_copy_invariants_after")
                if not isinstance(working_invariants, dict):
                    working_invariants = {}
                source_invariants_after = evidence.get("source_invariants_after")
                if not isinstance(source_invariants_after, dict):
                    source_invariants_after = {}
                if (
                    evidence.get("outcome")
                    in {"succeeded_append_only", "completed_with_warnings"}
                    and evidence.get("working_copy_invariants_unchanged") is True
                    and source_invariants_after.get("unchanged") is True
                    and isinstance(
                        working_invariants.get(
                            "trusted_review_status_identity_digest"
                        ),
                        dict,
                    )
                    and Path(evidence.get("prepared_output", "")).resolve()
                    == (folder / "ncs.db").resolve()
                ):
                    return "api-validation"
                checkpoint_path = folder / "api-checkpoint/api_checkpoint.json"
                if checkpoint_path.is_file() and report.get("parent_database") and report.get("sources"):
                    checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
                    if not isinstance(checkpoint, dict):
                        return None
                    checkpoint_baseline = checkpoint.get("baseline") or {}
                    checkpoint_identity = checkpoint.get("identity") or {}
                    if not isinstance(checkpoint_identity, dict):
                        return None
                    source_invariants = checkpoint_identity.get("source_invariants") or {}
                    if not isinstance(checkpoint_baseline, dict) or not isinstance(
                        source_invariants, dict
                    ):
                        return None
                    if (
                        isinstance(
                            checkpoint_baseline.get(
                                "trusted_review_status_identity_digest"
                            ),
                            dict,
                        )
                        and isinstance(
                            source_invariants.get(
                                "trusted_review_status_identity_digest"
                            ),
                            dict,
                        )
                    ):
                        return "api-collection"
            if report.get("kind") == "excel-delta":
                delta = json.loads((folder / "delta.json").read_text(encoding="utf-8"))
                if delta.get("ok") is True and self._excel_evidence_matches(folder, delta):
                    return "excel-validation"
        except (OSError, ValueError, TypeError, sqlite3.DatabaseError):
            return None
        return None

    def _excel_evidence_matches(self, folder: Path, evidence: dict) -> bool:
        expected = evidence.get("candidate_integrity")
        if not isinstance(expected, dict) or not isinstance(
            expected.get("trusted_review_status_identity_digest"), dict
        ):
            return False
        db = folder / "ncs.db"
        actual = {
            "sha256": file_sha256(db),
            "raw_ksa_sha256": raw_ksa_sha256(db),
            "trusted_review_status_identity_digest": trusted_review_status_identity_digest(db),
        }
        return (
            Path(evidence.get("candidate_db", "")).resolve() == db.resolve()
            and all(actual[key] == expected.get(key) for key in actual)
            and file_sha256(db) == actual["sha256"]
        )

    def resume(self, version: str) -> dict:
        with self.exclusive("resume", version) as builder_context:
            folder = self._version_dir(version)
            kind = self.resume_kind(version)
            if not kind:
                raise BuilderError(
                    "재개할 완료 기록이 없습니다. 이 버전을 완료로 사용하지 않습니다."
                )
            report = json.loads((folder / "build.json").read_text(encoding="utf-8"))
            previous_operation = report.get("operation_lineage")
            if previous_operation:
                report.setdefault("operation_history", []).append(previous_operation)
            report["operation_lineage"] = builder_context.lineage()
            report.update(
                status="building", resumed_at=datetime.now(timezone.utc).isoformat()
            )
            self._write_version_json(folder, "build.json", report, builder_context,
                                     action="resume", version=version)
            try:
                if kind == "api-collection":
                    require_builder_context(
                        builder_context, action="resume", root=self.root, state_dir=self.state,
                        version=version, version_dir=folder,
                    )
                    evidence = refresh_ncs_api_evidence(
                        Path(report["parent_database"]),
                        sources=report["sources"],
                        apply=True,
                        output_path=folder / "ncs.db",
                        checkpoint_dir=folder / "api-checkpoint",
                        resume=True,
                        builder_context=builder_context,
                        progress=self.progress,
                    )
                    self._write_version_json(folder, "api-refresh.json", evidence, builder_context,
                                             action="resume", version=version)
                    if evidence.get("outcome") not in {
                        "succeeded_append_only",
                        "completed_with_warnings",
                    }:
                        raise BuilderError(api_failure_message(evidence))
                if kind.startswith("api-"):
                    evidence = json.loads(
                        (folder / "api-refresh.json").read_text(encoding="utf-8")
                    )
                    self.progress(
                        "저장된 API 결과 재사용: 원문·검토 상태 보존 확인 (재수집 없음)"
                    )
                    expected = evidence.get("working_copy_invariants_after") or {}
                    expected_identity = expected.get(
                        "trusted_review_status_identity_digest"
                    )
                    if not isinstance(expected_identity, dict):
                        raise BuilderError(
                            "저장된 API 검토 상태 identity digest가 없어 안전한 재개를 중단했습니다. "
                            "새 API 수집을 시작하세요."
                        )
                    if raw_ksa_sha256(folder / "ncs.db") != expected.get(
                        "raw_ksa_sha256"
                    ) or trusted_review_status_counts(
                        folder / "ncs.db"
                    ) != expected.get("trusted_review_status_counts") or (
                        trusted_review_status_identity_digest(folder / "ncs.db")
                        != expected_identity
                    ):
                        raise BuilderError(
                            "저장된 API 결과와 후보 DB의 원문·검토 상태가 달라 재개를 중단했습니다."
                        )
                    report.update(
                        api_outcome=evidence["outcome"], sources=evidence["sources"]
                    )
                    if report.get("parent_database"):
                        parent_path = Path(report["parent_database"]).resolve()
                        parent_report = parent_path.parent / "build.json"
                        if (
                            parent_path.is_relative_to(self.state / "versions")
                            and parent_report.is_file()
                        ):
                            parent = json.loads(
                                parent_report.read_text(encoding="utf-8")
                            )
                            for key in ("source_delta", "ontology_processing"):
                                if key in parent:
                                    report[key] = parent[key]
                else:
                    evidence = json.loads(
                        (folder / "delta.json").read_text(encoding="utf-8")
                    )
                    if not self._excel_evidence_matches(folder, evidence):
                        raise BuilderError("저장된 Excel 결과와 후보 DB 무결성 증거가 달라 재개를 중단했습니다.")
                    report.update(
                        source_delta=evidence.get("source_delta"),
                        ontology_processing=evidence.get("ontology_processing"),
                    )
                self._write_version_json(folder, "build.json", report, builder_context,
                                         action="resume", version=version)
                return self._finish(folder, report, builder_context)
            except Exception as exc:
                self._failed(folder, report, exc, builder_context)
                raise

    def copy_current(self) -> dict:
        with self.exclusive("copy_current") as builder_context:
            source = self.current_db()
            self._space(source.stat().st_size * 2 + 1024**3)
            folder, report = self._new("current-copy", builder_context)
            try:
                self.progress(
                    "현재 MCP DB의 일관된 복사본 생성 (기존 API·검토 데이터 포함)"
                )
                _sqlite_online_snapshot(source, folder / "ncs.db", authorize=lambda: require_builder_context(
                    builder_context, action="copy_current", version_dir=folder))
                report["parent_database"] = str(source)
                return self._finish(folder, report, builder_context)
            except Exception as exc:
                self._failed(folder, report, exc, builder_context)
                raise
